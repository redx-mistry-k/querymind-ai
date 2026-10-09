"""Run the quota-conscious live Gemini SQL verification suite.

This script intentionally refuses to call Gemini unless the operator supplies
--confirm-free-tier after checking the API project's tier in Google AI Studio.
"""

from __future__ import annotations

import argparse
import math
import sqlite3
import sys
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from backend.app.core.config import Settings  # noqa: E402
from backend.app.main import create_app  # noqa: E402
from frontend.components.charts import build_figure  # noqa: E402
from frontend.components.exports import rows_to_csv, rows_to_xlsx  # noqa: E402

CASES = [
    (
        "Show the top 10 customers by revenue.",
        """SELECT c.name AS customer,
        ROUND(SUM(oi.quantity * oi.unit_price), 2) AS revenue
        FROM customers c JOIN orders o ON o.customer_id = c.id
        JOIN order_items oi ON oi.order_id = o.id
        WHERE o.status = 'completed'
        GROUP BY c.id, c.name ORDER BY revenue DESC LIMIT 10""",
        ("customer", "name"),
        ("revenue", "sales", "amount", "spent"),
        "series",
    ),
    (
        "Which product category generates the highest revenue?",
        """SELECT c.name AS category,
        ROUND(SUM(oi.quantity * oi.unit_price), 2) AS revenue
        FROM categories c JOIN products p ON p.category_id = c.id
        JOIN order_items oi ON oi.product_id = p.id
        JOIN orders o ON o.id = oi.order_id
        WHERE o.status = 'completed'
        GROUP BY c.id, c.name ORDER BY revenue DESC LIMIT 1""",
        ("category", "name"),
        ("revenue", "sales", "amount"),
        "first",
    ),
    (
        "Show monthly sales trends.",
        """SELECT strftime('%Y-%m', o.order_date) AS month,
        ROUND(SUM(oi.quantity * oi.unit_price), 2) AS sales
        FROM orders o JOIN order_items oi ON oi.order_id = o.id
        WHERE o.status = 'completed'
        GROUP BY month ORDER BY month""",
        ("month", "date", "period"),
        ("sales", "revenue", "amount"),
        "series",
    ),
    (
        "What is the average order value?",
        """WITH order_totals AS (
        SELECT o.id, SUM(oi.quantity * oi.unit_price) AS order_total
        FROM orders o JOIN order_items oi ON oi.order_id = o.id
        WHERE o.status = 'completed' GROUP BY o.id)
        SELECT ROUND(AVG(order_total), 2) AS average_order_value FROM order_totals""",
        (),
        ("average", "avg", "value", "amount"),
        "scalar",
    ),
    (
        "Which region has the highest number of orders?",
        """SELECT r.name AS region, COUNT(DISTINCT o.id) AS order_count
        FROM regions r JOIN customers c ON c.region_id = r.id
        JOIN orders o ON o.customer_id = c.id
        GROUP BY r.id, r.name ORDER BY order_count DESC LIMIT 1""",
        ("region", "name"),
        ("order_count", "orders", "count", "total"),
        "first",
    ),
]


def read_rows(path: Path, sql: str) -> list[dict[str, Any]]:
    connection = sqlite3.connect(path)
    connection.row_factory = sqlite3.Row
    try:
        return [dict(row) for row in connection.execute(sql).fetchall()]
    finally:
        connection.close()


def find_column(rows: list[dict[str, Any]], candidates: tuple[str, ...], numeric: bool) -> str:
    if not rows:
        raise AssertionError("The query returned no rows")
    for name, value in rows[0].items():
        lowered = name.lower()
        if any(token in lowered for token in candidates) and (
            isinstance(value, (int, float)) == numeric
        ):
            return name
    for name, value in rows[0].items():
        if isinstance(value, (int, float)) == numeric:
            return name
    raise AssertionError(f"Could not identify a {'numeric' if numeric else 'label'} column")


def verify_accuracy(
    actual: list[dict[str, Any]],
    expected: list[dict[str, Any]],
    label_tokens: tuple[str, ...],
    value_tokens: tuple[str, ...],
    mode: str,
) -> None:
    actual_value = find_column(actual, value_tokens, True)
    expected_value = find_column(expected, value_tokens, True)
    if mode == "scalar":
        assert math.isclose(
            float(actual[0][actual_value]),
            float(expected[0][expected_value]),
            rel_tol=1e-6,
            abs_tol=0.02,
        )
        return
    actual_label = find_column(actual, label_tokens, False)
    expected_label = find_column(expected, label_tokens, False)
    count = 1 if mode == "first" else len(expected)
    assert len(actual) >= count
    for got, wanted in zip(actual[:count], expected[:count], strict=True):
        assert str(got[actual_label]) == str(wanted[expected_label])
        assert math.isclose(
            float(got[actual_value]),
            float(wanted[expected_value]),
            rel_tol=1e-6,
            abs_tol=0.02,
        )


def run(model: str) -> None:
    base = Settings()
    if not base.gemini_api_key:
        raise SystemExit("GEMINI_API_KEY is not configured.")
    with tempfile.TemporaryDirectory(prefix="querymind-live-") as directory:
        root = Path(directory)
        settings = base.model_copy(
            update={
                "root_dir": root,
                "data_dir": root / "data",
                "demo_db_path": root / "data" / "demo" / "querymind_demo.db",
                "metadata_db_path": root / "data" / "querymind.db",
                "public_demo_mode": False,
                "api_access_token": "",
            }
        )
        with TestClient(create_app(settings)) as client:
            for index, (question, canonical_sql, labels, values, mode) in enumerate(CASES, 1):
                generated = client.post(
                    "/api/v1/query/generate",
                    json={
                        "question": question,
                        "database_id": "demo",
                        "provider": {
                            "provider": "gemini",
                            "model": model,
                            "temperature": 0,
                            "max_tokens": 768,
                        },
                        "conversation": [],
                    },
                )
                assert generated.status_code == 200, generated.text
                payload = generated.json()
                validation = payload["validation"]
                assert validation and validation["valid"], payload
                execution = client.post(
                    "/api/v1/query/execute",
                    json={
                        "validation_id": validation["validation_id"],
                        "question": question,
                        "provider": "gemini",
                        "model": model,
                    },
                )
                assert execution.status_code == 200, execution.text
                result = execution.json()
                expected = read_rows(settings.demo_db_path, canonical_sql)
                verify_accuracy(result["rows"], expected, labels, values, mode)
                chart_type = result["chart"]["type"]
                if chart_type not in {"table", "kpi"}:
                    frame = pd.DataFrame(result["rows"], columns=result["columns"])
                    figure = build_figure(frame, result["chart"], chart_type, question)
                    assert figure is not None and figure.to_json()
                assert result["summary"] and result["followups"]
                if index == 1:
                    assert rows_to_csv(result["rows"], result["columns"])
                    workbook = load_workbook(
                        BytesIO(rows_to_xlsx(result["rows"], result["columns"]))
                    )
                    assert workbook.active.max_row == result["row_count"] + 1
                    saved = client.post(
                        "/api/v1/queries/saved",
                        json={
                            "name": "Live Gemini verification",
                            "question": question,
                            "sql": payload["generation"]["sql"],
                            "database_id": "demo",
                            "category": "Verification",
                        },
                    )
                    assert saved.status_code == 201
                usage = payload.get("usage") or {}
                print(
                    f"{index}/5 ok | rows={result['row_count']} | chart={chart_type} "
                    f"| usage={usage}"
                )
            history = client.get("/api/v1/query/history").json()["history"]
            assert sum(item["status"] == "success" for item in history) == len(CASES)
            assert client.get("/api/v1/queries/saved").json()["saved_queries"]
    print(f"Live Gemini verification passed for {model}.")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--model", default="gemini-3.8-flash")
    parser.add_argument(
        "--confirm-free-tier",
        action="store_true",
        help="Confirm that this API project is shown as Free Tier in Google AI Studio.",
    )
    args = parser.parse_args()
    if not args.confirm_free_tier:
        parser.error(
            "Refusing live inference. Confirm the configured project's Free Tier status in "
            "Google AI Studio, then pass --confirm-free-tier."
        )
    run(args.model)


if __name__ == "__main__":
    main()
